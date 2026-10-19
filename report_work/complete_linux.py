from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent.parent
SRC=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（3.2重写稿）.docx'
DST=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（Linux补全稿）.docx'
d=Document(SRC)
templates={
 'body':deepcopy(next(p for p in d.paragraphs if p.text.startswith('本课题围绕实验室自研NPU'))._p.get_or_add_pPr()),
 'heading':deepcopy(next(p for p in d.paragraphs if p.text=='3.2.1 总体设计与部署方式' and p.style.name=='Heading 3')._p.get_or_add_pPr()),
 'caption':deepcopy(next(p for p in d.paragraphs if p.style.name=='Caption' and p.text.startswith('图1 '))._p.get_or_add_pPr())
}
# Reserve four figure numbers for the Linux section and update existing references.
for p in d.paragraphs:
 if re.search(r'图\d+',p.text):
  text=re.sub(r'图(\d+)',lambda m:'图'+str(int(m[1])+4),p.text)
  p.clear();p.add_run(text)

def boundary(title):
 ps=d.paragraphs;s=next(i for i,p in enumerate(ps) if p.text==title and p.style.name=='Heading 2')
 e=next(i for i in range(s+1,len(ps)) if ps[i].style.name in ['Heading 1','Heading 2'])
 anchor=ps[e]._p;node=ps[s]._p.getnext()
 while node is not anchor:
  nxt=node.getnext();node.getparent().remove(node);node=nxt
 return anchor
anchor=boundary('3.1 Linux系统相关')
def para(text,kind='body'):
 el=OxmlElement('w:p');el.append(deepcopy(templates[kind]));anchor.addprevious(el)
 p=Paragraph(el,d._body);p.add_run(text)
 for name in ['autoSpaceDE','autoSpaceDN']:
  pr=p._p.get_or_add_pPr();e=pr.find(qn('w:'+name))
  if e is None:e=OxmlElement('w:'+name);pr.append(e)
  e.set(qn('w:val'),'0')
 return p
def placeholder(num,title,description):
 p=para(f'图{num}待插入：{description}')
 p.paragraph_format.keep_with_next=True
 p.paragraph_format.keep_together=True
 p=para(f'图{num} {title}（图片待补充）','caption')
 p.paragraph_format.keep_with_next=False

para('3.1.1 系统总体方案','heading')
para('按照开题报告确定的研究方案，本阶段已完成自研RISC-VSoC上的Linux系统移植及驱动开发，建立了从处理器复位、引导程序执行、内核启动到用户态应用运行的完整软件环境。系统以FPGA实现的RISC-VSoC为硬件基础，围绕处理器、存储控制器、中断控制器、UART、显示与触摸模块以及NPU等设备完成软件适配，使操作系统能够管理硬件资源，并为应用程序提供统一的设备访问接口。')
para('Linux软件系统由引导程序、内核、设备树和根文件系统共同构成，各部分的组成及构建关系如图1所示。其中，引导程序负责准备内核启动环境；Linux内核提供进程管理、内存管理和设备管理等基础功能；设备树描述SoC的硬件资源；根文件系统提供命令行工具、运行库及应用程序。针对FPGA平台的资源和运行频率约束，系统采用定制引导程序、内核功能裁剪和精简根文件系统相结合的实现方式，在满足目标应用需求的基础上控制软件规模。')
placeholder(1,'Linux系统组成与构建流程','插入Linux系统组成与构建流程图。可复用开题报告中“Linux系统的组成及构建流程”图，展示引导程序、内核配置与源码、设备树、BusyBox及用户应用分别编译生成对应产物，再由Makefile统一组织、打包并下载到FPGA的过程。图中标明各模块的输入、输出及依赖关系。')

para('3.1.2 轻量级引导程序','heading')
para('引导程序是处理器复位后执行的基础软件，其主要任务是初始化关键硬件、建立必要的异常处理与固件服务，并将控制权交给Linux内核。针对本SoC的固定硬件组成，本阶段已完成轻量级定制引导程序的设计与实现，围绕硬件初始化、内存设置、内核加载及启动参数准备组织执行流程，减少与目标平台无关的功能配置。')
para('在启动早期，引导程序完成运行环境设置，并配置异常入口及必要的上下文保存、恢复机制；随后初始化UART，使启动过程能够通过串口输出调试信息，同时完成中断控制器、定时器等关键外设的基础设置。异常处理部分按事件类型组织分发和处理流程，为启动阶段的故障定位以及后续固件服务调用提供支持。图2给出了引导程序的执行过程。')
para('针对Linux内核与底层硬件之间的调用需求，已实现满足本系统启动和运行需要的最小化SBI服务，包括所需的控制台与定时器支持。内核通过约定接口请求底层服务，由固件侧完成相应处理。在启动交接阶段，引导程序将内核映像置于预定内存区域，准备设备树地址等启动信息，设置必要的异常委托与处理器运行状态，再转入Linux内核入口。通过上述实现，完成了从复位环境到操作系统运行环境的衔接。')
placeholder(2,'RISC-V引导程序执行流程','插入已实现的引导程序流程图，可参考开题报告中的对应图。主流程依次表示复位入口、基础运行环境设置、异常入口配置、UART及必要外设初始化、SBI服务准备、内核映像与设备树准备、异常委托与运行状态设置、进入Linux内核。旁支说明Linux通过SBI调用固件服务并返回的关系；具体寄存器名称和加载顺序以实现代码为准。')

para('3.1.3 内核裁剪与根文件系统','heading')
para('Linux内核移植围绕目标处理器能力、内存布局和外设组成进行配置。本阶段已完成RISC-V体系结构相关配置和平台资源适配，使内核能够识别目标SoC的处理器、内存及必要外设，并完成系统初始化。针对FPGA平台资源有限的特点，结合menuconfig等配置工具对内核功能进行裁剪，保留进程管理、虚拟内存管理、根文件系统支持及目标应用所需的驱动，移除未使用的设备、文件系统和其他可选功能。')
para('裁剪过程中，以系统能够完整启动和目标应用能够正常运行为基本条件，逐项检查配置依赖，避免因删除支撑模块导致驱动无法初始化或用户程序无法运行。调试阶段保留必要的串口日志及定位手段，在功能稳定后精简非必要的调试和分析配置。通过将内核配置与本SoC的硬件组成相对应，已形成适用于目标平台的定制内核镜像，为外设驱动和上层应用提供运行基础。')
para('根文件系统采用开题报告中确定的BusyBox结合手动构建方案。已完成BusyBox的交叉编译与功能配置，建立基本目录结构，并补充系统运行所需的库文件、设备节点、初始化脚本及测试应用。启动脚本负责组织用户空间初始化，准备必要的运行环境并进入命令行或应用程序，使内核启动后能够继续完成用户态功能。')
para('通过内核与根文件系统的配合，系统已能够从内核初始化进入用户空间，执行基本命令并运行外设测试程序。裁剪后的内核配置、BusyBox配置及根文件系统组成共同构成当前软件基线，便于在后续功能扩展时定位新增模块带来的空间和启动开销。镜像体积与启动时间的具体结果将在最终实验材料中按统一统计口径补充，分别说明内核映像、根文件系统和完整部署镜像的大小，以及启动计时的起止位置。')

para('3.1.4 设备树与外设驱动','heading')
para('设备树用于向Linux描述SoC的硬件组成，是内核与自研外设建立联系的重要依据。本阶段已完成处理器、内存、中断控制器、定时器及各外设的设备树适配，描述相应寄存器地址范围、中断资源和设备间的连接关系，并使其与硬件地址映射保持一致。驱动根据设备描述获取所需资源，完成设备识别、初始化和接口注册，将底层寄存器操作封装为应用程序能够使用的访问方式。')
para('对于UART等通用外设，采用内核现有驱动并完成平台配置；对于需要定制适配的中断控制器、显示与触摸模块以及NPU，则依据对应硬件接口完成驱动开发。驱动实现围绕资源获取、寄存器读写、运行状态管理和中断响应展开，并完成与内核构建配置及设备树的集成。通过分别检查设备匹配、资源地址和中断处理，建立了从硬件事件到内核响应的完整路径。')
para('图像显示部分已完成显示模块及相关DMA传输的驱动适配，按照硬件要求配置显示数据缓冲区、传输地址和控制状态，使应用程序生成的图像数据能够送至LCD显示。触摸输入部分完成输入信息读取与上报，并与显示应用结合实现基本人机交互。该过程需要保持图像格式、缓冲区布局和硬件读取方式一致，同时协调数据传输、状态更新和事件处理，确保显示与输入功能能够配合运行。')
para('NPU驱动完成了加速器资源管理与执行控制接口的适配，为上层程序提供参数和指令数据准备、地址配置、计算启动及状态获取等支持。通过驱动将NPU控制操作纳入Linux设备访问流程，应用程序能够在操作系统环境中组织处理器与加速器的协同工作。显示、触摸及NPU等外设的访问关系如图3所示，各驱动分别处理对应硬件细节，上层应用按功能调用相应接口。')
placeholder(3,'Linux设备驱动与应用协同关系','插入应用、Linux驱动和SoC外设的分层结构图。顶层为图形交互及加速器测试应用，中间层为显示与触摸驱动、NPU驱动及中断处理，底层为LCD、触摸模块、DMA、NPU和存储空间。用箭头分别表示控制请求、图像数据传输和完成事件；在旁侧标出设备树向驱动提供寄存器与中断资源的关系。驱动框架名称及具体调用接口按实际代码填写。')

para('3.1.5 系统集成与运行验证','heading')
para('在各模块完成后，已通过统一Makefile组织引导程序、Linux内核、设备树及根文件系统的编译与打包，形成可用于目标SoC的系统镜像。构建过程统一管理各部分的配置与产物依赖，将内核、设备描述和用户态环境组合为完整部署内容，并通过JTAG下载至FPGA平台进行系统联调。统一构建流程使配置修改能够传递到相应产物，便于重复生成和验证同一软件版本。')
para('系统验证按照启动过程和应用功能逐步展开。首先通过串口观察引导程序初始化、内核启动、设备识别和根文件系统进入情况，确认从处理器复位到用户空间的运行链路；随后运行基本命令及外设测试程序，验证UART、显示、触摸和NPU等设备的访问功能；最后通过应用程序组织外设协同工作，并在LCD上完成显示与交互演示。上述移植、驱动开发和整体联调工作均已完成，实现了开题报告所确定的Linux系统软件功能。')
placeholder(4,'Linux启动及外设应用运行结果','插入真实板端运行材料，建议组成三幅子图：（a）UART启动日志，保留引导程序输出、Linux内核初始化和进入用户空间的关键内容；（b）LCD显示及触摸交互实物照片；（c）Linux下NPU驱动测试或应用调用的运行记录，展示任务启动与完成结果。注明板卡型号、CPU工作频率、内核版本及测试日期；镜像大小和启动耗时由对应构建产物与日志记录填写。')
para('本阶段形成了轻量级引导程序、内核配置与镜像、设备树、定制外设驱动、BusyBox根文件系统以及统一构建脚本等软件成果。系统已具备基本命令行、图形交互和加速器访问能力，解决了自研硬件从裸机控制到操作系统管理的适配问题，并为后续AI编译产物部署、应用联调和性能测量提供了配套运行环境。')

anchor=boundary('4.1 Linux系统相关')
para('按照开题报告确定的研究内容，目前已完成RISC-VSoC上的Linux系统移植及驱动开发。软件系统已形成引导程序、Linux内核、设备树和根文件系统相互配合的完整结构，并在FPGA平台完成从复位启动到用户态应用运行的整体联调。系统能够通过Linux访问SoC外设，运行基本命令和测试程序，并支持LCD显示、触摸交互及NPU控制等应用功能，完成了本阶段的Linux软件建设任务。')
para('在启动环境方面，已实现面向目标SoC的轻量级定制引导程序，完成关键外设初始化、异常处理框架、必要SBI服务及内核启动交接。在操作系统方面，已完成与硬件能力相匹配的内核配置和功能裁剪，并基于BusyBox制作精简根文件系统。在设备支持方面，已完成设备树编写及相关外设驱动适配，将寄存器访问、中断响应和数据传输封装为Linux下的设备接口。这些工作共同解决了启动环境、内核资源描述和定制外设访问之间的衔接问题。')
para('与开题阶段的总体方案相比，Linux部分已由各模块的设计规划落实为能够构建、部署和运行的软件系统。通过统一Makefile管理编译与镜像集成，形成了可重复使用的软件基线；通过板端启动和应用演示，验证了各部分能够协同工作。后续将在现有成果基础上整理启动日志、系统配置、镜像组成及外设运行记录，补充统一口径的体积和时间数据，并继续为AI编译器的优化实验与系统应用验证提供支持。')

# Update the status cells that previously deferred all Linux progress reporting.
for table in d.tables:
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    if p.text=='具体进展在Linux章节补充':p.text='已完成，详见3.1节'

for style in d.styles:
 if re.fullmatch(r'toc [123]',style.name,re.I):
  style.paragraph_format.line_spacing=1.25
  style.paragraph_format.space_before=Pt(0)
  style.paragraph_format.space_after=Pt(0)
d.save(DST)
new=Document(DST)
text='\n'.join(p.text for p in new.paragraphs)
assert not re.search(r'[\u4e00-\u9fff][ \t]+[A-Za-z]|[A-Za-z][ \t]+[\u4e00-\u9fff]',text)
heads=[p.text for p in new.paragraphs if p.style.name=='Heading 3' and p.text.startswith('3.1.')]
assert len(heads)==5
caps=[p.text for p in new.paragraphs if p.style.name=='Caption' and re.match(r'图\d+',p.text)]
assert [int(re.match(r'图(\d+)',s)[1]) for s in caps]==list(range(1,12)),caps
print(DST)
print('\n'.join(heads))
print('Figure captions 1–11 verified; Chinese-English whitespace check passed.')
