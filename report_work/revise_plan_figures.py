from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent.parent
SRC=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（Linux补全稿）.docx'
DST=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（研究计划与示意图修订稿）.docx'
d=Document(SRC);ps=d.paragraphs
bodypr=deepcopy(next(p for p in ps if p.text.startswith('本课题围绕实验室自研NPU'))._p.get_or_add_pPr())
headpr=deepcopy(next(p for p in ps if p.style.name=='Heading 2' and p.text.startswith('5.1'))._p.get_or_add_pPr())
cappr=deepcopy(next(p for p in ps if p.style.name=='Caption' and p.text.startswith('图1 '))._p.get_or_add_pPr())

def picture(p,name):
 p.clear();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
 pf.line_spacing=1;pf.space_before=Pt(8);pf.space_after=Pt(3)
 pf.keep_together=True;pf.keep_with_next=True
 p.add_run().add_picture(str(ROOT/'report_work/new_figures'/f'{name}.png'),width=Cm(15.1))
for num,name in [(1,'linux_build'),(2,'linux_boot'),(3,'linux_drivers')]:
 p=next(p for p in d.paragraphs if p.text.startswith(f'图{num}待插入：'))
 picture(p,name)
 cap=next(p for p in d.paragraphs if p.style.name=='Caption' and p.text.startswith(f'图{num} '))
 cap.text=cap.text.replace('（图片待补充）','')

table3=deepcopy(d.tables[2]._tbl);table4=deepcopy(d.tables[3]._tbl)
start=next(p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith('5.'))
end=next(p for p in d.paragraphs if p.style.name=='Heading 1' and p.text.startswith('6.'))
node=start._p.getnext()
while node is not end._p:
 nxt=node.getnext();node.getparent().remove(node);node=nxt

def para(s='',kind='body'):
 el=OxmlElement('w:p');el.append(deepcopy({'body':bodypr,'heading':headpr,'caption':cappr}[kind]));end._p.addprevious(el)
 p=Paragraph(el,d._body)
 for part in re.split(r'(\{cite:\d+\})',s):
  if not part:continue
  if part.startswith('{cite:'):
   r=p.add_run('['+part[6:-1]+']');r.font.superscript=True;r.font.size=Pt(9)
  else:p.add_run(part)
 for name in ['autoSpaceDE','autoSpaceDN']:
  pr=p._p.get_or_add_pPr();e=pr.find(qn('w:'+name))
  if e is None:e=OxmlElement('w:'+name);pr.append(e)
  e.set(qn('w:val'),'0')
 return p
def fig(num,title,name):
 picture(para(),name)
 p=para(f'图{num} {title}','caption');p.paragraph_format.keep_with_next=False
def table(num,title,el):
 p=para(f'表{num} {title}','caption');p.paragraph_format.keep_with_next=True
 if num==4:
  # Keep this short schedule on one page with its caption.
  rows=el.findall(qn('w:tr'))
  for row in rows[:-1]:
   for pe in row.iter(qn('w:p')):
    pp=pe.find(qn('w:pPr'))
    if pp is None:pp=OxmlElement('w:pPr');pe.insert(0,pp)
    e=pp.find(qn('w:keepNext'))
    if e is None:e=OxmlElement('w:keepNext');pp.append(e)
    e.set(qn('w:val'),'1')
 end._p.addprevious(el)

para('在已完成Linux系统移植和NPU离线编译流程的基础上，后续研究将重点围绕内存规划优化与算子融合展开。前者通过分析中间张量的使用关系，减少推理过程所需的存储空间；后者结合宏指令和硬件数据通路，减少相邻运算之间的数据搬运与控制开销。两项工作均以当前可运行的YOLO网络为基线，通过单算子、典型子图和完整网络的实验验证其正确性与实际效果。')
para('5.1 内存规划优化','heading')
para('当前编译器主要按照执行顺序为中间结果分配空间，虽然便于组织地址，但已经完成使用的张量仍可能保留原有区域。后续拟在现有内存规划中加入生命周期分析，记录每个张量的生产节点、全部消费者、实际物理大小和对齐要求。当某个张量已完成全部消费后，其空间可供后续张量复用，从而降低运行区的峰值需求。图11以连续执行的三个算子说明这一过程，图中仅展示复用原理，不代表具体模型的实测收益。')
fig(11,'基于张量生命周期的内存复用示意','memory_reuse')
para('实现时将以现有顺序分配作为对照，首先在固定执行顺序下建立可用空间管理与分配机制，再结合目标NPU的4通道对齐规则检查候选区域是否满足大小和访问要求。外部输入、最终输出和模型参数按各自的保留需求管理，避免在仍有用途时被回收。对已采用预布局的通道拼接，还需将底层存储区域与各分支视图关联起来，只有相关消费者均完成访问后，才允许复用共享空间。')
para('在此基础上，将分支执行顺序与内存分配共同考虑。多分支网络在满足数据依赖的情况下可以具有不同的合法顺序，不同顺序会改变共享输入和分支输出的存活时间。MAGIS等研究说明执行调度与存储优化之间具有联合研究的空间{cite:6}。针对本平台，拟从当前可执行节点中选择候选操作，根据新增张量大小、可释放空间以及后续拼接需求进行评分，并通过有限范围搜索选择峰值占用较低的执行方案。')
para('该部分的研究重点是使生命周期复用、通道布局与拼接共享关系相互兼容。每个候选方案均需通过依赖顺序、地址范围和读写冲突检查，并设置固定顺序复用、仅调整分支顺序及联合优化等实验组，分析不同机制的实际贡献。评价时分别记录峰值内存、对齐填充、编译耗时和推理时间，以说明优化对部署空间与执行效率的不同影响。')

para('5.2 算子融合','heading')
para('现有宏指令能够完成较大粒度的张量运算，但相邻算子分别执行时，仍可能产生中间结果写回、重新读取和重复配置的开销。后续将优先分析网络中出现频繁的卷积与激活组合，研究在卷积输出阶段直接完成激活处理的方法。DNNFusion等工作表明，合理融合能够减少中间数据和执行开销{cite:7}。本课题拟结合自研NPU的数据通路和指令配置开展相应扩展，其基本方式如图12所示。')
fig(12,'卷积与激活融合的拟实现方式','operator_fusion')
para('具体实现将与硬件设计人员协同，明确融合所需的激活模式、输出位宽、量化移位及饱和处理顺序，并通过指令字段或配置寄存器表达相应控制信息。编译器在计算图中识别符合条件的连续算子，将其转换为融合执行描述，同步调整中间存储分配、参数组织及指令生成。对于硬件尚不支持的组合，或者激活前结果仍被其他节点使用的情况，保留原有执行方式，保证模型语义和数据依赖正确。')
para('融合研究将以减少实际数据搬运为目标，同时考察硬件修改带来的资源和时序代价。跨层执行中的存储访问成本需要结合硬件层次分析{cite:8}，因此将分别统计融合前后的指令数量、中间结果读写量、NPU执行周期，以及FPGA资源占用和可工作频率。在完成小规模组合验证后，再将融合路径应用于完整网络，判断其收益是否能够覆盖硬件扩展成本。')

para('5.3 实验验证与工作安排','heading')
para('实验首先保证优化前后的计算正确性。针对硬件采用的整数或定点运算，建立与权重、偏置、累加、移位和饱和规则一致的软件参考；量化位置和数值格式也是验证的重要依据{cite:9}。测试覆盖非4倍数通道、不同空间尺寸、多消费者张量及连续拼接等情况，并由单算子逐步扩展到子图和完整YOLO网络。只改变地址和执行顺序的优化应保持同一硬件基线的输出结果；融合若改变数值处理过程，则需明确误差来源及可接受的任务指标范围。')
para('性能实验以同一模型、输入和明确的软硬件配置开展对照，分别测量工作区峰值、指令数量、NPU执行周期，以及包含装载、预处理和后处理的端到端时间。通过重复测量记录波动范围，区分空间节省、数据搬运减少和工作频率变化所产生的影响。实验分组如表3所示，其中融合扩展组在硬件功能完成并通过正确性验证后开展。')
table(3,'编译优化对照实验设计',table3)
para('后续按“完善基线—实现优化—板端验证—论文整理”的顺序推进，优先完成生命周期复用及分支调度，再根据瓶颈分析推进算子融合。具体时间安排见表4。预期形成面向本NPU的内存规划方法、相应的融合编译支持，以及能够说明适用条件和优化效果的实验结果，并据此完成论文撰写。')
table(4,'中期检查后的研究安排',table4)

d.save(DST)
new=Document(DST)
assert len(new.tables)==4
heads=[p.text for p in new.paragraphs if p.style.name=='Heading 2' and re.match(r'5\.\d',p.text)]
assert len(heads)==3,heads
pending=[p.text for p in new.paragraphs if '待插入：' in p.text]
assert len(pending)==2 and pending[0].startswith('图4') and pending[1].startswith('图10')
captions=[int(re.match(r'图(\d+)',p.text)[1]) for p in new.paragraphs if p.style.name=='Caption' and re.match(r'图\d+',p.text)]
assert captions==list(range(1,13)),captions
assert not re.search(r'[\u4e00-\u9fff][ \t]+[A-Za-z]|[A-Za-z][ \t]+[\u4e00-\u9fff]','\n'.join(p.text for p in new.paragraphs))
print(DST);print('\n'.join(heads));print('5 diagrams inserted; 2 real-evidence placeholders retained; figure numbering 1–12 verified.')
