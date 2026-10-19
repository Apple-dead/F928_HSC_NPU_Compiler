from docx import Document
from pathlib import Path
from zipfile import ZipFile
from PIL import Image
root=Path(__file__).resolve().parent.parent
out=root/'report_work'/'npu_sources';out.mkdir(exist_ok=True)
for j,name in enumerate(['NPU编译器.docx','NPU编译器技术文档.docx']):
 d=Document(root/name); lines=[]
 for i,p in enumerate(d.paragraphs):
  blips=p._p.xpath('.//a:blip')
  imgs=[str(d.part.related_parts[b.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')].partname) for b in blips]
  if p.text or imgs: lines.append(f'{i} {p.style.name}: {p.text}'+(' IMAGES:'+str(imgs) if imgs else ''))
 for i,t in enumerate(d.tables): lines.append('TABLE '+str(i)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for r in t.rows))
 (out/(name+'.txt')).write_text('\n'.join(lines),encoding='utf-8')
 with ZipFile(root/name) as z:
  for m in z.namelist():
   if m.startswith('word/media/'):
    f=out/(str(j)+'_'+Path(m).name);f.write_bytes(z.read(m))
    try:
     im=Image.open(f);print(f.name,im.size)
    except: pass
 print(name,'paragraphs',len(d.paragraphs),'tables',len(d.tables))
d=Document(root/'论文中期材料/蒋一平-硕士论文中期检查报告（修订稿）.docx')
active=False
for i,p in enumerate(d.paragraphs):
 if p.text.startswith('3.2') and p.style.name.startswith('Heading'):active=True
 if active:print(i,p.style.name,p.text, 'DRAWING' if p._p.xpath('.//a:blip') else '')
 if active and p.text.startswith('4.') and p.style.name.startswith('Heading'):break
