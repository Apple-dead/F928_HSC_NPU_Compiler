from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent.parent
SRC=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（修订稿）.docx'
DST=ROOT/'论文中期材料/蒋一平-硕士论文中期检查报告（3.2重写稿）.docx'
d=Document(SRC);ps=d.paragraphs
start=next(p for p in ps if p.text=='3.2 AI编译器相关' and p.style.name=='Heading 2')
end=next(p for p in ps if p.text=='4.已研究工作总结' and p.style.name=='Heading 1')
bodytemplate=deepcopy(ps[69]._p.get_or_add_pPr())
headingtemplate=deepcopy(ps[68]._p.get_or_add_pPr())
captemplate=deepcopy(ps[73]._p.get_or_add_pPr())
saved={}
for num in [1,3,4]:
 idx=next(i for i,p in enumerate(ps) if p.style.name=='Caption' and p.text.startswith(f'图{num} '))
 saved[num]=deepcopy(ps[idx-1]._p)
table=deepcopy(d.tables[1]._tbl)
# Preserve everything outside the requested section, apart from necessary figure references.
outside_before=[p.text for p in ps[:ps.index(start)] if not p.style.name.startswith('toc')]
node=start._p.getnext()
while node is not end._p:
 nxt=node.getnext();node.getparent().remove(node);node=nxt

def para(text='',kind='body'):
 el=OxmlElement('w:p');end._p.addprevious(el)
 el.append(deepcopy({'body':bodytemplate,'heading':headingtemplate,'caption':captemplate}[kind]))
 p=Paragraph(el,d._body)
 for s in re.split(r'(\{cite:\d+\})',text):
  if not s:continue
  if s.startswith('{cite:'):
   r=p.add_run('['+s[6:-1]+']');r.font.superscript=True;r.font.size=Pt(9)
  else:p.add_run(s)
 for name in ['autoSpaceDE','autoSpaceDN']:
  pr=p._p.get_or_add_pPr();e=pr.find(qn('w:'+name))
  if e is None:e=OxmlElement('w:'+name);pr.append(e)
  e.set(qn('w:val'),'0')
 return p
def fig(num,caption,path=None,width=15.1):
 if path:
  p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  pf=p.paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(6);pf.space_after=Pt(0);pf.line_spacing=1;pf.keep_with_next=True
  p.add_run().add_picture(str(ROOT/'report_work/npu_sources'/path),width=Cm(width))
 else:end._p.addprevious(deepcopy(saved[num]))
 p=para(f'图{num} {caption}','caption');p.paragraph_format.keep_with_next=False

para('3.2.1 总体设计与部署方式','heading')
para('本课题围绕实验室自研NPU构建离线AI编译器，将神经网络模型转换为目标硬件能够直接执行的指令和参数文件。目前已完成模型解析、参数整理、内存规划、指令生成及文件导出等主要模块，并在集成NPU的RISC-VSoC上完成FPGA验证，能够运行编译器生成的YOLO网络二进制文件并正确进行图像识别。编译器的设计目标是将原先需要人工完成的模型信息提取、数据排列、地址计算和指令配置组织为统一流程，为自研NPU的模型部署与后续优化提供软件基础。')
para('编译系统采用PC侧编译、SoC侧加载执行的部署方式，如图1所示。在PC侧，编译器依据模型结构和目标硬件约束，预先确定运算顺序、参数布局及中间结果地址，生成完整的NPU指令流。部署时，将指令、参数和输入数据装载到SoC的指定存储区域，由RISC-V处理器完成必要的地址配置与启动控制，随后NPU按指令流读取数据、执行计算并写回结果。端侧不需要重新解析模型计算图，加载与控制程序主要负责连接编译产物和硬件执行接口。')
fig(1,'离线AI编译器与RISC-V异构SoC的关系')
para('编译器总体流程如图2所示，主要由模型前端、内存规划、算子指令生成和文件导出等部分组成。前端将模型结构转换为统一的中间表示（IR），并提取权重、偏置等参数；内存规划模块结合硬件约束确定计算过程及数据地址；后端据此生成汇编指令，经汇编编码后形成机器码，最终输出部署所需文件。各阶段通过模型结构文件和内存规划文件传递信息，使模型解析与硬件指令生成保持相对独立。')
para('图2给出了编译流程的各阶段及主要产物。其中，PT2是当前已实现的模型输入路径，ONNX表示可通过新增前端进行适配的扩展方向。图中的“NPU可执行文件”指目标NPU可加载的指令及参数产物；在BRAM初始化场景下可按布局合并组织，在DDR部署场景下则分别导出指令和参数二进制文件，以适配不同的装载地址。')
fig(2,'NPU编译器总体流程','1_image4.png',14.4)

para('3.2.2 模型解析与参数整理','heading')
para('模型解析的主要任务是从原始模型中提取计算结构与参数，并转换为后续模块能够统一处理的形式。目前前端以PyTorch导出的PT2模型为输入，读取计算节点、张量形状及参数映射关系，将算子类别、输入输出连接、卷积核大小、步长和填充等信息写入图中间表示，同时保存模型权重与偏置。统一的IR使后端能够依据模型语义完成处理，而不必在每个编译阶段重复读取原始模型。该分层方式与通过中间表示逐级完成模型到目标代码转换的思路一致{cite:3}。')
para('以连续的卷积和ReLU运算为例，前端首先记录两个节点的运算属性及数据依赖，明确ReLU的输入来自卷积输出；随后提取卷积所需的权重和偏置，供参数生成与地址规划使用。现有前端已建立卷积、激活、池化、最近邻上采样、填充、通道拼接和全连接等操作到内部表示的映射，具体可执行的参数范围由后续硬件适配规则进一步约束。若需要支持其他模型格式，可增加相应解析入口，将其转换为相同的IR，并完成算子语义和数值行为的校验后复用后端流程。')
para('模型参数还需按照NPU的数据读取方式进行整理。框架导出的权重排列与硬件使用的通道分组、像素交织顺序可能不同，因此参数生成模块需要完成权重重排、必要的通道补齐及偏置组织。对于采用4通道对齐的情况，编译器以通道组为单位安排存储，将不足一组的部分补齐，使指令所描述的通道数量与实际参数空间相匹配。图3给出了逻辑通道与物理存储之间的关系。')
fig(3,'通道对齐对物理存储空间的影响示意')
para('参数整理与执行计划共同决定硬件实际读取的数据。例如，输出通道较多的卷积会根据目标硬件能力拆分为多个计算分组，每个分组需要对应的权重范围和输出区域。编译器依据分组结果生成参数数据及偏移信息，保证后续指令能够访问正确的参数。对于具有偏置和不具有偏置的操作，也分别组织相应内容，避免参数长度或排列方式与指令配置不一致。')

para('3.2.3 内存规划','heading')
para('内存规划是连接模型描述与指令生成的关键环节。该阶段以统一IR和硬件配置为依据，确定输入数据、模型参数、中间张量、模型输出及指令所需的存储空间，并形成memory_plan.json等规划文件。规划结果不仅记录各区域的起始地址和大小，还包含算子分组、参数偏移及输出位置等执行信息，供参数生成、指令生成与文件导出共同使用。通过在编译阶段完成这些工作，可以减少人工计算地址的工作量，并保持各阶段对同一数据的解释一致。')
para('仍以卷积后接ReLU为例，规划器首先根据输入形状、卷积属性和通道对齐规则计算卷积输出的物理大小，为其分配存储区域；随后将该输出地址作为ReLU的输入地址，并为激活结果安排输出位置。与只按逻辑张量元素数量估计空间相比，实际地址计算还需要考虑对齐填充、分组方式和硬件寻址规则。规划完成后，编译器检查相关地址区间的边界及重叠情况，后端据此生成一致的寄存器配置。')
para('针对模型运算规模与NPU单次计算能力不一致的情况，内存规划同时组织相应的执行步骤。例如，卷积输出通道可按硬件支持的粒度进行分组，逻辑上的一次卷积由多次计算完成，各组分别使用对应的参数和输出区域。对于部分需要组合实现的运算，则依据目标指令能力安排相邻操作及数据连接。这一过程将计算图中的逻辑节点转换为可以实际执行的操作序列，使后续指令生成集中于具体字段的配置和编码。')
para('对于满足条件的通道拼接，当前编译器采用预先安排生产者输出位置的方法，使各分支直接写入最终拼接张量的对应区域，如图4所示。当拼接方向、空间尺寸及通道对齐满足目标布局要求时，后续算子可以直接读取该区域，从而省去单独的拼接数据拷贝。这一处理将网络中的连接关系与物理存储布局结合起来，减少了部分不必要的数据搬运。')
fig(4,'满足布局条件时的通道拼接免拷贝示意')
para('目前中间结果区域仍以顺序分配为主，尚未实现面向全图生命周期的通用缓冲区复用。随着网络深度和分支数量增加，已完成消费的中间张量仍可能占用原有空间。同时，拼接预布局会使多个逻辑张量共享同一物理存储对象的不同区域，复用时需要统一处理其别名和使用关系。因此，现有规划文件也为后续开展生命周期分析、分支调度与存储联合优化提供了基础。')

para('3.2.4 指令生成与文件导出','heading')
para('在模型执行顺序、数据布局和存储地址确定后，算子编译模块读取内存规划文件，为各执行节点生成对应的汇编片段。目标NPU采用宏指令形式，通常先通过配置指令设置输入地址、输出地址、参数地址和算子属性，再由计算指令触发相应硬件单元。编译器因此需要在指令生成前完成尺寸、通道分组和地址等静态计算，使硬件能够依据预先生成的配置执行运算。')
para('以卷积运算为例，后端先设置输入、输出和权重地址所在的通用寄存器，再配置卷积核大小、通道数量及偏置等相关字段，最后生成CONV指令启动计算；后续ReLU节点则引用卷积结果地址并生成相应配置与计算指令。图5给出了卷积与ReLU的汇编示例，展示了地址配置、算子属性配置和计算启动之间的关系。图中地址用于说明该示例的编译结果，实际模型的地址由对应内存规划决定。')
fig(5,'卷积与ReLU的汇编指令示例','1_image9.png',15.1)
para('汇编器按照目标指令编码规则，将上述汇编序列转换为32位机器码。构建脚本进一步组织参数与指令导出：对于BRAM仿真和初始化场景，可生成按预定布局组织的COE等文件；对于DDR部署场景，则分别输出npu_instr.bin与npu_params.bin，供系统按指定基地址加载。指令编码、参数排列和装载位置均与内存规划保持对应，构成从模型到硬件执行文件的完整转换过程。')
para('分层的后端结构也为增加硬件算子提供了接口。新增算子时，需要在IR中补充其输入输出和属性，在规划阶段确定形状、空间需求及合法参数范围，再增加相应汇编生成规则；若引入新的指令编码，还需同步扩展汇编器。由此可在保留整体流程的基础上开展软硬件协同扩展，并通过单算子与完整网络验证检查各阶段的一致性。')

para('3.2.5 系统验证与阶段成果','heading')
para('本阶段已将编译器与实际硬件执行过程贯通。NPU已集成到RISC-VSoC中，并随整个SoC部署到FPGA；系统能够加载本编译器生成的YOLO网络二进制文件，由NPU执行网络计算并正确完成图像识别。这一结果表明，模型解析、参数整理、内存规划、指令编码及板端装载之间已建立可运行的完整链路，为后续在同一硬件平台上比较不同编译策略提供了实验基础。')
para('目前已确认的是完整系统中的功能验证结果。为进一步形成可重复、可比较的实验记录，后续将补充板卡型号、工作频率、模型版本、输入尺寸及量化配置等条件，并记录NPU执行周期和端到端处理时间。现阶段尚未形成统一的帧率、功耗、加速比或数据集精度统计，相关信息按表2继续整理。')
para('表2 FPGA验证记录与待补充实验信息','caption').paragraph_format.keep_with_next=True
end._p.addprevious(table)
para('图6待插入：RISC-VSoC的FPGA实物与YOLO运行结果。左侧放置FPGA开发板实物照片，标注板卡型号及图像输入、结果输出接口；右侧放置同次运行的输入图像与真实识别结果，并注明模型名称、输入尺寸和测试日期。可在下方补充指令、参数装载及执行完成的真实日志。')
para('图6 FPGA系统运行与YOLO图像识别结果（实测材料待补充）','caption')
para('通过上述工作，已解决自研NPU部署过程中模型信息、数据布局和指令格式之间的转换问题，形成可重复构建的离线编译流程。后续研究将在此基础上完善逐层数值比对与性能测量，重点分析中间张量的生命周期、分支执行次序和布局约束对存储需求的影响，进一步实现执行调度与存储联合优化，并通过对照实验评价其效果。')

# Shift the later optimization figure reference after adding the assembly example.
after=False
for p in d.paragraphs:
 if p._p is end._p:after=True
 if after and '图6' in p.text:
  for r in p.runs:r.text=r.text.replace('图6','图7')
d.save(DST)
new=Document(DST)
newtext='\n'.join(p.text for p in new.paragraphs)
assert not re.search(r'[\u4e00-\u9fff][ \t]+[A-Za-z]|[A-Za-z][ \t]+[\u4e00-\u9fff]',newtext)
heads=[p.text for p in new.paragraphs if p.style.name=='Heading 3' and p.text.startswith('3.2.')]
assert len(heads)==5,heads
assert len(d.tables)==len(Document(SRC).tables)
print(DST)
print('\n'.join(heads))

