from pathlib import Path
import re,json,math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from refs import REFS

ROOT=Path(__file__).resolve().parent.parent
WORK=ROOT/'report_work'; OUT=ROOT/'论文中期材料'; OUT.mkdir(exist_ok=True)
FIG=WORK/'figures'; FIG.mkdir(exist_ok=True)
WIDTH=8640 # 15.24 cm usable width
# Preset: academic paper; named override: Beihang supplied proposal and midterm rules.
TOKENS={'page_cm':[21,29.7],'margins_cm':[2.54,2.54,3.17,2.59],
 'body':'宋体 / Times New Roman, 12pt, 1.5 lines, first line 24pt',
 'h1':'黑体, 15pt, before12pt, after6pt','h2':'黑体,14pt, before9pt, after3pt',
 'h3':'宋体 bold,12pt,before6pt,after3pt','caption':'10.5pt centered',
 'table_width_dxa':WIDTH,'table_cell_margin_dxa':80,'table_font_pt':10.5,
 'header_footer_cm':[1.25,1.25],'reference':'10.5pt,1.2 lines,hanging21pt'}
(WORK/'design_tokens.json').write_text(json.dumps(TOKENS,ensure_ascii=False,indent=2),encoding='utf8')

def draw_fig(name,kind):
 W=1600
 heights={'taxonomy':690,'system':730,'pipeline':720,'layout':600,'concat':600,'optimize':680}
 im=Image.new('RGB',(W,heights[kind]),'white'); d=ImageDraw.Draw(im)
 font=lambda n:ImageFont.truetype('C:/Windows/Fonts/simsun.ttc',n)
 def txt(x,y,s,n=38,anchor='mm',fill='black'):
  d.multiline_text((x,y),s,font=font(n),fill=fill,anchor=anchor,align='center',spacing=12)
 def box(x,y,w,h,s,n=38,fill='#F6F7F8'):
  d.rounded_rectangle((x,y,x+w,y+h),radius=6,fill=fill,outline='#242424',width=3)
  txt(x+w/2,y+h/2,s,n)
 def arrow(x1,y1,x2,y2):
  d.line((x1,y1,x2,y2),fill='black',width=3)
  a=math.atan2(y2-y1,x2-x1)
  points=[(x2,y2),(x2-17*math.cos(a-.45),y2-17*math.sin(a-.45)),(x2-17*math.cos(a+.45),y2-17*math.sin(a+.45))]
  d.polygon(points,fill='black')
 if kind=='taxonomy':
  labs=['模型与图语义\n形状、类型、依赖','图与存储优化\n重写、融合、复用','硬件映射\n分块、布局、调度','目标产物\n代码或NPU命令流']
  for i,s in enumerate(labs):
   box(35+i*395,70,345,160,s,35)
   if i<3: arrow(380+i*395,150,425+i*395,150)
  box(250,320,1100,130,'目标约束：指令能力、存储容量、对齐、数值格式',38)
  for x in [600,990]: arrow(x,320,x,240)
  box(380,535,840,105,'设备侧：装载、启动、同步与结果读取',38)
  arrow(1400,235,1400,580);arrow(1400,580,1230,580)
 elif kind=='system':
  box(45,160,410,240,'PC侧离线编译\n模型解析与优化\n指令／参数生成',40)
  d.rectangle((630,50,1555,660),outline='black',width=3)
  txt(1090,90,'FPGA：已集成NPU的RISC-VSoC',39)
  box(710,160,340,150,'RISC-V处理器\n加载与执行控制',37)
  box(1150,160,330,150,'NPU\n神经网络计算',39)
  box(790,450,610,135,'存储空间\n输入、参数、中间结果与输出',37)
  arrow(455,235,700,235);txt(570,180,'二进制',34)
  arrow(1050,235,1140,235)
  arrow(880,310,880,440);arrow(1300,310,1300,440)
  txt(800,705,'系统结构示意；实际总线与加载接口以平台实现为准',32)
 elif kind=='pipeline':
  labs=['PT2模型\n图与参数','图中间表示\nmodel_ir.json','静态执行与存储计划\nmemory_plan.json','目标指令生成\n汇编与32位编码']
  for i,s in enumerate(labs):
   box(35+i*395,60,345,165,s,32)
   if i<3:arrow(380+i*395,142,425+i*395,142)
  box(200,350,470,140,'参数转换\n权重、偏置与物理布局',37)
  box(930,350,470,140,'COE与二进制导出\n指令镜像＋参数镜像',37)
  arrow(990,225,990,290);arrow(990,290,435,290);arrow(435,290,435,340)
  arrow(1420,225,1420,320);arrow(1420,320,1170,320);arrow(1170,320,1170,340)
  arrow(670,420,920,420)
  box(480,590,640,95,'RISC-VSoC加载，NPU执行',39)
  arrow(1170,490,1170,550);arrow(1170,550,800,550);arrow(800,550,800,580)
 elif kind=='layout':
  txt(800,45,'示例：单个空间位置的6个有效通道按4通道对齐',38)
  txt(175,160,'逻辑通道',36)
  for i in range(6):box(360+i*145,110,135,100,'C'+str(i),36,'white')
  arrow(800,220,800,300)
  txt(175,380,'物理通道',36)
  for i in range(8):box(340+i*145,330,135,100,'C'+str(i) if i<6 else '补齐',36,'#EEEEEE' if i>=6 else 'white')
  d.line((340,465,910,465),fill='black',width=3);d.line((920,465,1490,465),fill='black',width=3)
  txt(620,510,'第1通道组',36);txt(1200,510,'第2通道组',36)
  txt(800,575,'实际占用包含有效数据与对齐填充；图示不表示实测层大小',30)
 elif kind=='concat':
  box(60,70,500,130,'分支A生产者\n满足布局约束',39)
  box(60,330,500,130,'分支B生产者\n满足布局约束',39)
  d.rectangle((880,65,1510,465),outline='black',width=3)
  box(900,85,590,150,'拼接输出的A区域',39,'white')
  box(900,255,590,190,'拼接输出的B区域',39,'#EEEEEE')
  arrow(560,135,890,160);arrow(560,395,890,350)
  txt(720,90,'直接写入',34);txt(720,440,'直接写入',34)
  txt(1190,510,'同一物理存储对象的不同区域',34)
  txt(800,575,'无额外拼接拷贝；复用时必须共同检查别名与消费者',34)
 elif kind=='optimize':
  box(45,70,430,145,'图依赖与静态形状\n现有NPU布局约束',38)
  box(585,70,430,145,'合法调度候选\n分支顺序与布局方案',38)
  box(1125,70,430,145,'存储对象及视图\n生命周期与地址分配',36)
  arrow(475,140,575,140);arrow(1015,140,1115,140)
  box(1090,350,465,145,'代价估计与合法性检查\n峰值、对齐、读写冲突',35)
  arrow(1340,215,1340,340)
  box(45,350,430,145,'选择方案并生成指令\n板端验证与消融实验',36)
  arrow(1090,420,485,420)
  arrow(1090,380,800,380);arrow(800,380,800,225)
  txt(800,330,'有限搜索与反馈',33)
  txt(800,615,'后续拟实现方法；当前已完成基线编译与FPGA功能验证',33)
 im.save(FIG/(name+'.png'),dpi=(300,300))
for x in ['taxonomy','system','pipeline','layout','concat','optimize']:draw_fig(x,x)

TABLES={
'plan':('表1 论文研究工作计划', [2.8,5.8,6.64],[
 ['时间安排','计划内容','阶段说明'],
 ['2025.12—2026.02','构建RISC-VSoC与FPGA验证环境，集成相关外设','原计划中的硬件平台阶段'],
 ['2026.03—2026.04','研究启动加载与基础软件适配','系统软件方向'],
 ['2026.05—2026.07','Linux内核与根文件系统移植','具体进展在Linux章节补充'],
 ['2026.08—2026.09','外设驱动开发与系统联调','具体进展在Linux章节补充'],
 ['2026.10—2026.12','面向NPU的AI编译与优化研究','已有全流程基线，继续开展方法与实验'],
 ['2027.01—2027.04','补充实验、撰写与修改学位论文','优化验证与论文整理交叉进行'],
 ['2027.05','论文答辩','按学院实际安排执行']]),
'evidence':('表2 FPGA验证记录与待补充实验信息',[4.0,5.62,5.62],[
 ['项目','当前状态','需归档的证据'],
 ['系统集成与运行','已完成RISC-VSoC集成NPU及FPGA运行','系统版本、板卡照片、运行日志'],
 ['YOLO图像识别','已运行编译器生成的二进制并正确识别','真实输入与输出图像、模型及样本说明'],
 ['硬件与模型配置','待补充完整记录','FPGA型号、频率、模型版本、输入尺寸、量化格式'],
 ['数值一致性','待补充系统性测试结果','软件参考、逐层输出、误差指标'],
 ['存储与性能数据','待测量并汇总','工作区峰值、执行周期、计时边界及重复次数']]),
'experiments':('表3 编译优化对照实验设计',[3.8,5.9,5.54],[
 ['实验组','配置','主要比较目的'],
 ['基线组','固定顺序＋现有顺序分配','建立功能、存储与周期基线'],
 ['生命周期组','固定顺序＋别名安全的空间复用','分析静态复用带来的容量收益'],
 ['调度组','改变合法分支顺序，保持统一分配策略','分析调度对生命周期和峰值的影响'],
 ['联合优化组','分支调度＋布局约束＋空间复用','评价联合决策相对独立阶段的收益'],
 ['融合扩展组（条件性）','加入经验证的硬件后处理融合','比较访存、周期、频率与硬件资源'],
 ['共同验证','相同模型、输入、频率和数值约定','逐层正确性、任务表现与可复现性']]),
'future':('表4 中期检查后的研究安排',[3.1,7.1,5.04],[
 ['时间','主要任务','阶段产物'],
 ['2026.10—2026.11','整理硬件约束与基线；实现生命周期和别名分析','能力说明、数值测试、存储分析'],
 ['2026.12—2027.01','实现布局约束下的复用及分支联合调度','算法原型、合法性检查、对照结果'],
 ['2027.02—2027.03','完成多负载测试及消融；按条件评估硬件扩展','FPGA实测数据与适用范围分析'],
 ['2027.03—2027.04','整理实验，完成论文写作与修改','论文初稿、修改稿与答辩材料'],
 ['2027.05','完成答辩及材料归档','按学院安排完成提交']]),
'related':('表1 代表性研究与宏指令NPU编译的关系',[3.05,5.7,6.49],[
 ['研究方向','代表工作及主要方法','对本课题的启示与适用边界'],
 ['编译表示','TVM、MLIR、TensorIR：分层表示与目标映射','保留图语义和硬件约束；按需求控制实现规模'],
 ['图结构优化','TASO、DNNFusion：图替换与融合','需核对定点语义；软件分组不等于消除硬件访存'],
 ['自动调度','Ansor、ROLLER：搜索或构造有效程序','仅优化硬件实际暴露的分组、布局与调度参数'],
 ['MCU部署','DORY、MCUNet：分块、存储与执行协同','借鉴资源建模；DMA、计算内核和目标硬件不同'],
 ['跨层存储','MCUNetV2、StreamNet、MAGIS：分块、流式与联合规划','综合考虑空间、重算和别名；不直接搬用训练场景结论'],
 ['异构编译','Deeploy、MATCH：异构映射与静态代码生成','目标与本课题接近；C代码与NPU指令镜像形式不同'],
 ['硬件代价','MAESTRO、DeFiNES：数据流及层次化访问建模','先明确存储和执行能力，再评估跨层或指令扩展']])
}

def xmlval(parent,tag,attrs):
 e=parent.find(qn(tag))
 if e is None:e=OxmlElement(tag);parent.append(e)
 for k,v in attrs.items():e.set(qn(k),str(v))
 return e
def fontset(font,east='宋体',size=12,bold=False):
 font.name='Times New Roman';font.size=Pt(size);font.bold=bold;font.color.rgb=RGBColor(0,0,0)
 rp=font._element
 xmlval(rp,'w:rFonts',{'w:ascii':'Times New Roman','w:hAnsi':'Times New Roman','w:eastAsia':east})
def noautospace(p):
 pr=p._p.get_or_add_pPr()
 for tag in ['w:autoSpaceDE','w:autoSpaceDN']:xmlval(pr,tag,{'w:val':0})
def setup(doc):
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(2.54);sec.bottom_margin=Cm(2.54);sec.left_margin=Cm(3.17);sec.right_margin=Cm(2.59)
 sec.header_distance=Cm(1.25);sec.footer_distance=Cm(1.25)
 s=doc.styles['Normal'];fontset(s.font)
 p=s.paragraph_format;p.line_spacing=1.5;p.first_line_indent=Pt(24);p.space_before=Pt(0);p.space_after=Pt(0);p.widow_control=True
 p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 xmlval(s.element.get_or_add_pPr(),'w:autoSpaceDE',{'w:val':0});xmlval(s.element.get_or_add_pPr(),'w:autoSpaceDN',{'w:val':0})
 for i,size in [(1,15),(2,14),(3,12)]:
  s=doc.styles['Heading '+str(i)];fontset(s.font,'黑体' if i<3 else '宋体',size,i==3)
  p=s.paragraph_format;p.first_line_indent=Pt(0);p.left_indent=Pt(0);p.space_before=Pt({1:12,2:9,3:6}[i]);p.space_after=Pt({1:6,2:3,3:3}[i]);p.line_spacing=1.5;p.keep_with_next=True;p.keep_together=True;p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 for nm in ['Caption','Header','Footer']:
  s=doc.styles[nm];fontset(s.font,size=10.5);p=s.paragraph_format;p.first_line_indent=Pt(0);p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for nm in ['TOC 1','TOC 2','TOC 3']:
  if nm not in doc.styles:doc.styles.add_style(nm,1)
  s=doc.styles[nm];fontset(s.font,size=11);p=s.paragraph_format;p.first_line_indent=Pt(0);p.left_indent=Pt({'TOC 1':0,'TOC 2':14,'TOC 3':28}[nm]);p.line_spacing=1.15;p.space_before=Pt(0);p.space_after=Pt(0)
 doc.core_properties.author='蒋一平';doc.core_properties.subject='硕士论文中期材料';doc.core_properties.language='zh-CN'
def para(doc,text='',style=None,center=False,indent=True):
 p=doc.add_paragraph(style=style);noautospace(p)
 if not indent:p.paragraph_format.first_line_indent=Pt(0)
 if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if text:p.add_run(text)
 return p
def spacer(doc,pt=10.5):
 p=para(doc,indent=False);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(pt);p.add_run().font.size=Pt(1)
 return p
def field(p,code,display=''):
 r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
 r=p.add_run();x=OxmlElement('w:instrText');x.set(qn('xml:space'),'preserve');x.text=code;r._r.append(x)
 r=p.add_run();x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),'separate');r._r.append(x)
 p.add_run(display)
 r=p.add_run();x=OxmlElement('w:fldChar');x.set(qn('w:fldCharType'),'end');r._r.append(x)
def newsection(doc,start=1,fmt='decimal'):
 sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
 sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
 hp=sec.header.paragraphs[0];hp.text='北京航空航天大学电子信息工程学院';hp.style=doc.styles['Header'];noautospace(hp)
 border=OxmlElement('w:pBdr');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'4');bt.set(qn('w:color'),'000000');bt.set(qn('w:space'),'6');border.append(bt);hp._p.get_or_add_pPr().append(border)
 fp=sec.footer.paragraphs[0];fp.style=doc.styles['Footer'];field(fp,' PAGE ','1')
 xmlval(sec._sectPr,'w:pgNumType',{'w:start':start,'w:fmt':fmt})
 return sec
def cover(doc,kind):
 p=para(doc,'中图分类号：TN492',indent=False);p.paragraph_format.line_spacing=1.5
 p=para(doc,'学号：ZY2402113',indent=False)
 spacer(doc,42)
 p=para(doc,center=True,indent=False);p.add_run().add_picture(str(WORK/'original_media/蒋一平-开题报告/image1.png'),width=Cm(14.4))
 spacer(doc,26)
 p=para(doc,center=True,indent=False);p.paragraph_format.line_spacing=1
 r=p.add_run('硕士中期检查报告' if kind=='midterm' else '专业硕士学位文献综述');fontset(r.font,'华文行楷',40 if kind=='midterm' else 32)
 spacer(doc,62)
 p=para(doc,center=True,indent=False);p.paragraph_format.line_spacing=1.35
 title='基于RISC-V的\n软硬件协同及编译' if kind=='midterm' else '面向端侧NPU的\nAI编译器研究综述'
 r=p.add_run(title);fontset(r.font,size=32)
 spacer(doc,75)
 for s in ['作者姓名  蒋一平','学科专业  电子信息','指导教师  李洪革教授，李丽副教授','培养院系  电子信息工程学院']:
  p=para(doc,s,center=True,indent=False);p.paragraph_format.line_spacing=1.65
  for r in p.runs:fontset(r.font,size=14)
 spacer(doc,34)
 p=para(doc,'2026年10月',center=True,indent=False)
 doc.core_properties.title=title.replace('\n','')
def heading(doc,text,level):
 p=para(doc,text,'Heading '+str(level),indent=False)
 return p
def caption(doc,text):
 p=para(doc,text,'Caption',indent=False);p.paragraph_format.keep_with_next=False
 spacer(doc,10.5)
def figure(doc,name,cap):
 p=para(doc,center=True,indent=False);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=True
 r=p.add_run();shape=r.add_picture(str(FIG/(name+'.png')),width=Cm(15.0));shape._inline.docPr.set('descr',cap)
 caption(doc,cap)
def table(doc,key):
 cap,widths,rows=TABLES[key];spacer(doc,10.5)
 p=para(doc,cap,'Caption',indent=False);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(5)
 t=doc.add_table(rows=0,cols=len(widths));t.autofit=False
 # exact geometry; fixed width and matching column/cell widths
 ws=[round(WIDTH*w/sum(widths)) for w in widths];ws[-1]=WIDTH-sum(ws[:-1])
 pr=t._tbl.tblPr;xmlval(pr,'w:tblW',{'w:w':WIDTH,'w:type':'dxa'});xmlval(pr,'w:tblInd',{'w:w':80,'w:type':'dxa'});xmlval(pr,'w:tblLayout',{'w:type':'fixed'})
 mar=xmlval(pr,'w:tblCellMar',{})
 for k in ['top','bottom','left','right']:xmlval(mar,'w:'+k,{'w:w':80,'w:type':'dxa'})
 borders=xmlval(pr,'w:tblBorders',{})
 for k in ['top','left','bottom','right','insideH','insideV']:xmlval(borders,'w:'+k,{'w:val':'single','w:sz':4,'w:color':'777777'})
 for gc,w in zip(t._tbl.tblGrid.gridCol_lst,ws):gc.set(qn('w:w'),str(w))
 for ri,row in enumerate(rows):
  cells=t.add_row().cells;trpr=t.rows[-1]._tr.get_or_add_trPr();xmlval(trpr,'w:cantSplit',{})
  if ri==0:xmlval(trpr,'w:tblHeader',{})
  for ci,(cell,s,w) in enumerate(zip(cells,row,ws)):
   cell.width=Pt(w/20);xmlval(cell._tc.get_or_add_tcPr(),'w:tcW',{'w:w':w,'w:type':'dxa'})
   xmlval(cell._tc.get_or_add_tcPr(),'w:vAlign',{'w:val':'center'})
   if ri==0:xmlval(cell._tc.get_or_add_tcPr(),'w:shd',{'w:fill':'F0F0F0'})
   p=cell.paragraphs[0];noautospace(p);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci==0 else WD_ALIGN_PARAGRAPH.LEFT
   p.paragraph_format.keep_with_next=(ri==0)
   r=p.add_run(s);fontset(r.font,size=10.5,bold=ri==0)
 spacer(doc,10.5)
 if key=='plan' and hasattr(doc,'add_comment'):
  doc.add_comment(t.rows[5].cells[0].paragraphs[0].runs[0],text='原开题报告此项写作“2026.10~2027.12”，与其后2027年论文及答辩安排冲突。本初稿按2026年10—12月整理，请核对实际计划；不表示新增工作已经完成。',author='编写核对',initials='核对')
def citepara(doc,text,used):
 p=para(doc)
 for token in re.split(r'(\{cite:[^}]+\})',text):
  if token.startswith('{cite:'):
   for key in token[6:-1].split(','):
    if key not in used:used.append(key)
    r=p.add_run('['+str(used.index(key)+1)+']');r.font.superscript=True;r.font.size=Pt(9)
  else:p.add_run(token)
 return p
def bibliography(doc,used):
 for i,key in enumerate(used,1):
  p=para(doc,f'[{i}] '+REFS[key][0],indent=False)
  p.paragraph_format.left_indent=Pt(21);p.paragraph_format.first_line_indent=Pt(-21);p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_together=True
  for r in p.runs:fontset(r.font,size=10.5)
def toc(doc,level=2):
 p=para(doc,'目录',center=True,indent=False);fontset(p.runs[0].font,'黑体',15);p.paragraph_format.space_after=Pt(12)
 p=para(doc,indent=False);field(p,' TOC \\o "1-'+str(level)+'" \\h \\z \\u ','目录将在排版阶段更新。')
def build(kind):
 doc=Document();setup(doc);cover(doc,kind)
 text=(WORK/(kind+'.txt')).read_text(encoding='utf8')
 used=[];state='body'
 if kind=='literature':
  front,body=text.split('# 1.引言',1)
  newsection(doc,1,'upperRoman')
  for line in front.splitlines():
   if line=='@ABSTRACT':
    p=para(doc,'摘要',center=True,indent=False);fontset(p.runs[0].font,'黑体',15);p.paragraph_format.space_after=Pt(12)
   elif line=='@KEYWORDS':state='keywords'
   elif line=='@ENABSTRACT':
    doc.add_page_break();p=para(doc,'Abstract',center=True,indent=False);fontset(p.runs[0].font,size=15,bold=True);p.paragraph_format.space_after=Pt(12);state='en'
   elif line=='@ENKEYWORDS':state='enkeys'
   elif line:
    if state=='keywords':spacer(doc,12);para(doc,'关键词：'+line,indent=False)
    elif state=='enkeys':spacer(doc,12);para(doc,'Keywords: '+line,indent=False)
    else:para(doc,line)
  doc.add_page_break();toc(doc)
  newsection(doc,1);text='# 1.引言'+body
 else:
  newsection(doc,1,'upperRoman');toc(doc,3);newsection(doc,1)
 for line in text.splitlines():
  if not line:continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));p=heading(doc,line[level:].strip(),level)
   if line in ['# 4.结论','# 5.主要参考文献','# 7.参考文献']:p.paragraph_format.page_break_before=True
  elif line.startswith('@FIG '):
   name,cap=line[5:].split('|');figure(doc,name,cap)
  elif line.startswith('@TABLE '):table(doc,line[7:])
  elif line=='@BLANK':spacer(doc,18)
  elif line.startswith('@PLACEHOLDER '):
   p=para(doc,line[13:],indent=False);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
   for r in p.runs:fontset(r.font,size=10.5)
   b=xmlval(p._p.get_or_add_pPr(),'w:pBdr',{})
   for edge in ['top','left','bottom','right']:xmlval(b,'w:'+edge,{'w:val':'single','w:sz':4,'w:space':6,'w:color':'777777'})
  elif line.startswith('@CAPTION '):caption(doc,line[9:])
  elif line=='@REFS':bibliography(doc,used)
  elif line=='@GRADE':
   p=heading(doc,'附录 文献来源与等级核对',1);p.paragraph_format.page_break_before=True
   para(doc,'本综述共引用20篇正式发表的学术文献。来源等级统一按《中国计算机学会推荐国际学术会议和期刊目录（2022年）》核对，其中18篇来自A类、2篇来自B类；不将arXiv预印本单独作为等级依据，也不混用未注明版本的SCI分区。以下为检索核对信息，文献著录以正式出版记录为准。')
   groups={}
   for i,key in enumerate(used,1):
    _,venue,grade=REFS[key];groups.setdefault((venue,grade),[]).append(str(i))
   rows=[['正式发表来源','CCF等级（2022版）','本文文献序号']]
   for (v,g),ns in groups.items():rows.append([v,g,', '.join('['+n+']' for n in ns)])
   TABLES['grade']=('附表1 引用文献发表来源核对',[5,4,6.24],rows);table(doc,'grade')
   p=para(doc,'目录来源：https://www.ccf.org.cn/Academic_Evaluation/By_category/。出版信息核对渠道包括ACM、IEEE、USENIX与NeurIPS正式论文集、出版DOI及作者机构公开版本。检索整理日期：2026年10月16日。',indent=False)
   for r in p.runs:fontset(r.font,size=10.5)
  else:citepara(doc,line,used)
 # Prevent automatic Latin/Chinese spacing, including headers, cell styles and normal text.
 for p in doc.paragraphs:noautospace(p)
 name='蒋一平-硕士论文中期检查报告（初稿）.docx' if kind=='midterm' else '蒋一平-AI编译器文献综述.docx'
 path=OUT/name;doc.save(path)
 zh=len(re.findall('[\u4e00-\u9fff]',text));print(kind,'Chinese chars:',zh,'refs:',len(used),'paragraphs:',len(doc.paragraphs),'saved:',str(path))
 (WORK/(kind+'_citations.json')).write_text(json.dumps(used),encoding='utf8')
for kind in ['literature','midterm']:build(kind)
