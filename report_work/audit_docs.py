from pathlib import Path
from docx import Document
from lxml import etree
import re,json
root=Path(__file__).resolve().parent.parent
for f in (root/'论文中期材料').glob('*.docx'):
 d=Document(f);ps=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
 text='\n'.join(p.text for p in ps)
 bad=re.findall(r'[\u4e00-\u9fff][ \t]+[A-Za-z]|[A-Za-z][ \t]+[\u4e00-\u9fff]',text)
 print(f.name,'chars',len(re.findall('[\u4e00-\u9fff]',text)),'bad_spaces',bad[:20],'tables',len(d.tables),'figures',len(d.inline_shapes))
 print('pending Toc', '目录将在排版阶段更新' in text)
 refs=[]
 for p in d.paragraphs:
  if re.match(r'^\[\d+\]',p.text): refs.append(p.text)
 print('references',len(refs),'superscripts',sum(bool(r.font.superscript) for p in d.paragraphs for r in p.runs))
 if '中期' in f.name:
  for i,p in enumerate(d.paragraphs):
   if p.text in ['3.1 Linux系统相关','4.1 Linux系统相关']:print('linux empty',not d.paragraphs[i+1].text)
 # Check exact widths after WPS roundtrip.
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 for ti,t in enumerate(d.tables):
  grid=[int(x.get('{'+ns['w']+'}w')) for x in t._tbl.xpath('./w:tblGrid/w:gridCol')]
  target=int(t._tbl.xpath('./w:tblPr/w:tblW')[0].get('{'+ns['w']+'}w'))
  if sum(grid)!=target:print('WIDTH MISMATCH',ti,target,grid)
 (root/'report_work'/(('midterm' if '中期' in f.name else 'literature')+'_finaltext.txt')).write_text(text,encoding='utf8')
